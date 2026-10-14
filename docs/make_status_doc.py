"""
make_status_doc.py - Word document for the team and the guide: what was done phase by phase, the mathematical model
and its effect on accuracy, and the guide's BRI model with our changes. Numbers come from docs/final_results.json.
"""
import json
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
R = json.loads((ROOT / "docs" / "final_results.json").read_text(encoding="utf-8"))
FIG = ROOT / "docs" / "final_figs"
OUT = Path(r"C:\Users\ASUS\Desktop\Major_Project\docs\Project_Status_and_Mathematical_Model_2026-10-14.docx")
GREEN = "3E4A3F"

f8, f4 = R["fusion8"], R["fusion4"]
F8, F4 = f8["test"]["fused"], f4["test"]["fused"]
s8, s4 = f8["stats_fused_minus_best_single"], f4["stats_fused_minus_best_single"]
single4 = f4["test"]["singles"][f4["best_single_on_val"]]
bri = R["bri"]; enh = R["enh"]; st = R["stats4"]; wi = R["wiou"]
bri_single = json.loads((ROOT / "eval" / "bri" / "bri_validation_yolo11s_4cls_balanced_1024.json").read_text())

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(2.2))
sec.top_margin = sec.bottom_margin = Cm(2.0)
st_n = doc.styles["Normal"]; st_n.font.name = "Calibri"; st_n.font.size = Pt(11)
st_n.paragraph_format.space_after = Pt(6)
for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
    h = doc.styles[f"Heading {lvl}"]; h.font.name = "Calibri"; h.font.size = Pt(size)
    h.font.color.rgb = RGBColor(0x3E, 0x4A, 0x3F); h.font.bold = True


def para(text, bold_lead=None, style=None, italic=False, size=None, align=None):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True
        if size: r.font.size = Pt(size)
    r = p.add_run(text); r.italic = italic
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    return p


def bullet(text, bold_lead=None):
    return para(text, bold_lead, style="List Bullet")


def number(text, bold_lead=None):
    return para(text, bold_lead, style="List Number")


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_); tcPr.append(sh)


def table(rows, widths, bold_last=False, size=9.5, highlight=()):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(str(val)); r.font.size = Pt(size)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            if j > 0 and i > 0:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); shade(c, GREEN)
            elif (bold_last and i == len(rows) - 1) or i in highlight:
                r.bold = True; shade(c, "E8F0E3")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def formula(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.italic = True; r.font.name = "Cambria Math"; r.font.size = Pt(11.5)
    return p


def figure(path, width_cm, caption):
    doc.add_picture(str(path), width=Cm(width_cm)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(caption, italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)


def p1(x):
    return f"{x * 100:.1f}"


# ======================================================================= title
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("AI-Based Wood Quality Assessment"); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = RGBColor(0x3E, 0x4A, 0x3F)
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Project status, phase by phase, and the mathematical model for accuracy"); r.font.size = Pt(13)
para("Major Project, Dept. of CSBS, Dayananda Sagar College of Engineering, AY 2026-27 · Guide: Mrs. Aparna H D",
     align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
para("Team 4: Gurugubelli Gayatri (1DS23CB014), Prakyat Pramod Shetty (1DS23CB033), Saiprasad Rao (1DS23CB041), "
     "Yasti Anil Kotak (1DS23CB063) · 14 October 2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)

# ======================================================================= 1 summary
doc.add_heading("1. Summary", 1)
para("The project is complete. Everything proposed at the 50 % review has been implemented and measured on 12 physical "
     "boards that were never used for training or for any choice of settings. The final review is on 24 October 2026.")
para("", "At the 50 % review ")
doc.paragraphs[-1].add_run("both detectors were stuck at 63-65 mAP50 over 8 defect classes, and the larger model "
                           "did not help. The guide suggested a mathematical model to improve accuracy.")
para("", "Since then ")
doc.paragraphs[-1].add_run(f"we built that model (class balancing, weighted box fusion of detectors, Wise-IoU loss, per-class "
                           f"thresholds and statistical validation). The deployed detector raises test mAP50 from "
                           f"{p1(single4['map50'])} to {p1(F4['map50'])} on the four main defect classes, a statistically "
                           f"significant gain. We also implemented the guide's Board Recoverability Index (BRI) for grading, "
                           f"adapted it to our images and validated it.")
table([["Item", "Result on the 12 unseen test boards"],
       ["Detection, 8 classes (50 % review)", "YOLO11s 64.6 mAP50, YOLO11m 63.0 mAP50"],
       ["Best single model, 4 main classes", f"{p1(single4['map50'])} mAP50, {p1(single4['map'])} mAP50-95"],
       ["Deployed detector (fusion of four 4-class YOLO11s models)", f"{p1(F4['map50'])} mAP50, {p1(F4['map'])} mAP50-95"],
       ["Gain from the mathematical model (95 % interval)", f"+{s4['map50_diff_ci'][0]:.1f} to +{s4['map50_diff_ci'][1]:.1f} mAP50; better on {s4['boards_better']} of 12 boards (Wilcoxon p = {s4['wilcoxon_p']:.3f})"],
       ["Grading (BRI) vs grades from the annotations", f"weighted kappa {bri['test_per_image']['qwk']:.2f}; {bri['test_per_board']['exact']*100:.0f} % of boards graded exactly"],
       ["Application", "Web app: defects, heatmap, grade A / B / C / Reject, recoverable area, value estimate"]],
      [7.5, 9.0], highlight=(3, 4))

# ======================================================================= 2 phase 0
doc.add_heading("2. Phase 0: the starting point (before September)", 1)
para("The team's first model was YOLO11s trained on square 640-px images for 83 epochs on a laptop CPU. It reported "
     "62.2 mAP50 on its test set. Checking that pipeline before building on it showed three problems:")
number("The data was split image by image. Each physical board is cut into about 49 images, so the same 81 boards "
       "appeared in training, validation and test. The test score measured memory of boards, not generalisation.", "Board leakage. ")
number("The class names in the configuration (Crack, Knot, Hole, Stain, Split, Warp, Decay, Scratch) were not the "
       "dataset's classes. The real classes are Quartzity, Live_Knot, Marrow, Resin, Dead_Knot, Knot_with_crack, "
       "Knot_missing and Crack.", "Wrong class names. ")
number("The images are 2800 x 1024 strips. Squeezing them into a 640 x 640 square wasted 63 % of the input on padding "
       "and shrank thin defects such as cracks.", "Square input. ")
para("So the 62.2 was not a usable baseline, and the first task was to rebuild the data pipeline.")

# ======================================================================= 3 phase 1
doc.add_heading("3. Phase 1: up to the 50 % review (16 September)", 1)
doc.add_heading("3.1 What we did", 2)
bullet("4,000 line-scan images of sawn timber (Kodytek et al., 2022), 8 classes, 9,053 boxes after removing 158 "
       "zero-size labels.", "Dataset. ")
bullet("Split by physical board: 52 boards for training, 17 for validation, 12 for test, zero overlap (checked by the script).", "Leakage-free split. ")
bullet("Rectangular 1024-px input (1024 x 384 effective), no padding; augmentation that respects the grain direction "
       "(no rotation, no vertical flip).", "Input. ")
bullet("YOLO11m (20.0 M parameters) and YOLO11s (9.4 M), fine-tuned from COCO with one identical recipe on a Kaggle "
       "Tesla T4 GPU. Each model was evaluated on the test boards exactly once.", "Models. ")
bullet("A web application showing boxes, a heatmap, per-class statistics and a JSON report.", "Demo. ")
doc.add_heading("3.2 Results and the accuracy problem", 2)
T8 = R["test8"]
rows = [["Model (8 classes)", "Precision", "Recall", "F1", "mAP50", "mAP50-95", "ms / image (T4)"]]
for n, ms in (("YOLO11m", "33.2"), ("YOLO11s", "11.7")):
    a = T8[n]["ALL"]; rows.append([n] + [f"{x:.1f}" for x in a] + [ms])
table(rows, [3.6, 2, 2, 1.6, 1.8, 2.2, 2.6])
rows = [["Class", "Test boxes", "YOLO11m mAP50", "YOLO11s mAP50"]]
for c, nb in (("Dead_Knot", 343), ("Marrow", 46), ("Live_Knot", 406), ("Knot_missing", 24), ("Knot_with_crack", 30),
              ("Resin", 69), ("Crack", 95), ("Quartzity", 10)):
    rows.append([c, nb, f"{T8['YOLO11m'][c][3]:.1f}", f"{T8['YOLO11s'][c][3]:.1f}"])
table(rows, [4, 2.5, 3, 3])
para("The two models were tied, so a bigger network was not the answer. The score was held down by the data: "
     "Quartzity (125 training boxes) was not learned at all (13-15 mAP50), Crack is thin and faint (53-58), and the "
     "rare knot types were confused with each other. The training set is 37:1 imbalanced (2,899 Live_Knot boxes against "
     "78 Knot_missing). This is the accuracy issue that led the guide to suggest a mathematical model.")

# ======================================================================= 4 phase 2
doc.add_heading("4. Phase 2: from 50 % to completion (17 September - 8 October)", 1)
para("Every item promised at the 50 % review was implemented. Each choice was made on the 17 validation boards and each "
     "candidate final model was tested once, so the test numbers are honest.")
table([["Promised at 50 %", "What we did", "Outcome"],
       ["More samples for rare classes", "Kept the 4 most frequent classes; copy-paste balancing of real defect crops", "Better on validation, not significant on test"],
       ["Higher-resolution input", "Trained at 1280 px (1280 x 480) as well as 1024 px", "Not better on validation"],
       ["Mathematical model: fusion, Wise-IoU, thresholds, statistics", "Implemented all four (Section 5)", "Fusion adopted: significant gain"],
       ["CWB-style architecture", "Wise-IoU (the CWB loss) implemented and tested", "Not better; CondConv / BiFormer left as future work"],
       ["Image enhancement (CLAHE, denoising)", "Study under 6 imaging conditions on validation", "No gain on normal images (Section 4.3)"],
       ["Recoverability grading rule", "Guide's BRI implemented, adapted and validated (Section 6)", "Weighted kappa 0.83"],
       ["Market-price analytics", "Grade-to-value estimate in the application", "Done (inputs editable)"],
       ["Multi-seed runs", "Final setup repeated with 3 seeds", f"Spread {wi['seed_spread']:.1f} mAP50"],
       ["Final evaluation and application", "Single-shot test; app with grading and value", "Done; final deck prepared"]],
      [4.5, 6.5, 5.5], size=9)
doc.add_heading("4.1 Why four classes", 2)
para("The final system detects the four defect classes with the most data: Live_Knot, Dead_Knot, Resin and Crack "
     "(4,018 / 2,896 / 646 / 512 boxes). Crack, the most severe defect, is kept. The other four classes have too few "
     "examples (78-125 training boxes) to learn or to measure reliably on 12 test boards. On these four classes the "
     "8-class YOLO11s scores 69.6 mAP50 on test; that, not 64.6, is the fair starting point for every 4-class comparison.")
doc.add_heading("4.2 All runs", 2)
v4 = R["val4"]; t4 = R["test4"]
rows = [["Model (YOLO11s unless stated)", "Val mAP50", "Test mAP50", "Test mAP50-95"],
        ["8-class YOLO11s, scored on the 4 classes", "73.8", "69.6", "32.9"]]
for k, lab in (("plain_1024", "4 classes, no balancing, 1024 px"), ("full_1024", "4 classes, full balance 1:1"),
               ("sqrt_1024", "4 classes, partial balance 3:1, 1024 px (best single)"), ("sqrt_1280", "4 classes, partial balance, 1280 px"),
               ("plain_1280", "4 classes, no balancing, 1280 px")):
    rows.append([lab, f"{v4[k]['ALL'][3]:.1f}", f"{t4[k]['ALL'][3]:.1f}" if t4.get(k) else "not tested",
                 f"{t4[k]['ALL'][4]:.1f}" if t4.get(k) else "-"])
wiou_v = [v for k, v in wi["val_map50"].items() if "wiou" in k][0]
rows.append(["Wise-IoU loss, final setup", f"{wiou_v:.1f}", "not tested", "-"])
rows.append(["Seeds 42 / 1 / 2, final setup", " / ".join(f"{x:.1f}" for x in wi["seed_val_map50"]), "-", "-"])
rows.append(["Deployed: fusion of the four 4-class models", p1(f4["chosen"]["val_map50"]), p1(F4["map50"]), p1(F4["map"])])
table(rows, [7.5, 2.8, 2.8, 3], highlight=(len(rows) - 1,))
para("Models without a test number were rejected on validation by rules written down before training; testing them "
     "anyway would have meant choosing on the test set. Test evaluations in total: 7.", italic=True, size=9.5)
doc.add_heading("4.3 Image-enhancement study (synopsis hypothesis H1)", 2)
rows = [["Condition (validation mAP50)", "Raw", "CLAHE", "Denoise", "Sharpen"]]
for c, lab in (("normal", "Normal"), ("low_light", "Low light"), ("overexposed", "Over-exposed"), ("blur", "Blur"),
               ("noise", "Noise"), ("low_contrast", "Low contrast")):
    rows.append([lab] + [p1(enh["table"][c][e]["map50"]) for e in ("raw", "clahe", "denoise", "sharpen")])
table(rows, [5, 2.4, 2.4, 2.4, 2.4])
para("CLAHE lowered accuracy significantly in 5 of 6 conditions; denoising helped only on noisy images (+3 to +7 "
     "points). Hypothesis H1 is therefore rejected for normal images. Lighting matters far more than enhancement: low "
     "light costs about 25 points.")

# ======================================================================= 5 math model
doc.add_heading("5. The mathematical model for accuracy", 1)
para("This section is the main research contribution. It answers the guide's question: can a mathematical model "
     "improve accuracy, and by how much? Yes: by about 5.5 mAP50 points on the main defect classes, with a 95 % "
     "confidence interval that excludes zero.")
doc.add_heading("5.1 What it is built from", 2)
para("", "1. Class-balancing rule. ")
doc.paragraphs[-1].add_run("Rare classes are given more training examples by pasting real defect crops onto defect-free "
                           "training boards (inside the wood, along the grain, training split only). How many is set by a formula, "
                           "the repeat-factor rule of LVIS (Gupta et al., 2019), related to the effective number of samples of "
                           "Cui et al. (2019):")
formula("n_c* = √(n_c × n_max)        (partial balance; full balance: n_c* = n_max)")
para("It turns the 9.5 : 1 imbalance of the four classes into 3.1 : 1 (partial) or 1.06 : 1 (full).")
para("", "2. Weighted box fusion (Solovyev et al., 2021). ")
doc.paragraphs[-1].add_run("Several detectors look at the same board; boxes of the same class that overlap are merged into "
                           "one box, averaged by confidence, and a box that only one model sees is down-weighted:")
formula("b_f = Σ w_m c_i b_i / Σ w_m c_i        c_f = (Σ w_m c_i / Σ w_m) × min(Σ w_m, W) / W")
para("b = box, c = confidence, w_m = weight of model m, W = sum of all model weights.", size=9.5, italic=True)
para("", "3. Wise-IoU v3 box loss (Tong et al., 2023). ")
doc.paragraphs[-1].add_run("The loss used by the CWB-YOLOv8 reference model, which focuses training on ordinary-quality boxes:")
formula("L = r × exp(d² / (W_g² + H_g²)) × (1 − IoU),    r = β / (δ α^(β − δ))")
para("", "4. Per-class confidence thresholds. ")
doc.paragraphs[-1].add_run("One threshold per class, chosen to maximise F1 on validation.")
para("", "5. Statistical validation. ")
doc.paragraphs[-1].add_run("Board-level bootstrap (the 12 test boards resampled 1,000 times) for a 95 % confidence interval "
                           "of every improvement, and an exact Wilcoxon signed-rank test on the per-board differences.")
doc.add_heading("5.2 Did it improve the score?", 2)
rows = [["Step (test, 4 classes)", "mAP50", "mAP50-95", "Verdict"],
        ["8-class YOLO11s (50 % review), on the 4 classes", "69.6", "32.9", "starting point"],
        ["4-class model, no balancing", f"{t4['plain_1024']['ALL'][3]:.1f}", f"{t4['plain_1024']['ALL'][4]:.1f}", "not significant"],
        ["+ partial class balancing (best single on validation)", f"{p1(single4['map50'])}", f"{p1(single4['map'])}", "not significant"],
        ["+ 1280-px model", f"{t4['sqrt_1280']['ALL'][3]:.1f}", f"{t4['sqrt_1280']['ALL'][4]:.1f}", "not chosen on validation"],
        ["Weighted box fusion of the four 4-class models", p1(F4["map50"]), p1(F4["map"]), "significant gain"]]
table(rows, [7.5, 2.2, 2.4, 4.3], highlight=(5,))
table([["Fusion vs best single model (test)", "Difference", "95 % interval", "Boards better / worse", "Wilcoxon p"],
       ["4 classes, mAP50", f"+{(F4['map50'] - single4['map50']) * 100:.1f}", f"+{s4['map50_diff_ci'][0]:.1f} to +{s4['map50_diff_ci'][1]:.1f}", f"{s4['boards_better']} / {s4['boards_worse']}", f"{s4['wilcoxon_p']:.3f}"],
       ["4 classes, mAP50-95", f"+{(F4['map'] - single4['map']) * 100:.1f}", f"+{s4['map_diff_ci'][0]:.1f} to +{s4['map_diff_ci'][1]:.1f}", "-", "-"],
       ["8 classes (YOLO11m + YOLO11s), mAP50", f"+{(F8['map50'] - f8['test']['singles'][f8['best_single_on_val']]['map50']) * 100:.1f}", f"{s8['map50_diff_ci'][0]:+.1f} to {s8['map50_diff_ci'][1]:+.1f}", f"{s8['boards_better']} / {s8['boards_worse']}", f"{s8['wilcoxon_p']:.2f}"],
       ["8 classes, mAP50-95", f"+{(F8['map'] - f8['test']['singles'][f8['best_single_on_val']]['map']) * 100:.1f}", f"+{s8['map_diff_ci'][0]:.1f} to +{s8['map_diff_ci'][1]:.1f}", "-", "-"]],
      [5.5, 2.2, 3.4, 3, 2.2])
pc = F4["ap50"]; sd = t4["sqrt_1024"]
rows = [["Class (test mAP50)", "Best single model", "Deployed fusion", "Change"]]
for c in ("Dead_Knot", "Live_Knot", "Resin", "Crack"):
    rows.append([c, f"{sd[c][3]:.1f}", p1(pc[c]), f"{pc[c] * 100 - sd[c][3]:+.1f}"])
table(rows, [5, 3.5, 3.5, 2.5])
para("Fusion helps every class, most of all Crack (+11.5) and Resin (+5.5), the two thinner and rarer defects.")
doc.add_heading("5.3 What did not help (also a finding)", 2)
bullet("helped on validation (+2.2) but not on test; full balance was worst, because only 160 cracks could be cut out "
       "cleanly and each was pasted up to 17 times.", "Class balancing alone ")
bullet("was not better on validation.", "1280-px input alone ")
bullet(f"scored {wiou_v:.1f} vs 75.3 on validation, within the seed spread of {wi['seed_spread']:.1f}: not adopted.", "Wise-IoU ")
bullet(f"lowered macro-F1 on test ({p1(f4['thresholds']['fused_test_per_class']['F1'])} vs {p1(f4['thresholds']['fused_test_global']['F1'])} with one global threshold): they over-fit the few validation boxes.", "Per-class thresholds ")
para("Fusion works because the four models make different mistakes: different input sizes and different class "
     "balance. Averaging them cancels some errors and places boxes more precisely. Its cost is speed: four "
     "models run per image, about 4x the single-model time (about 2 seconds per image on a laptop CPU).")
doc.add_heading("5.4 Benefits and limits", 2)
bullet("a measured, significant accuracy gain with a confidence interval, and every choice made before seeing the test set.", "Benefits: ")
bullet("12 test boards is a small test set; severity weights and grade thresholds still need an expert; the four rare classes are out of scope.", "Limits: ")

# ======================================================================= 6 BRI
doc.add_heading("6. The guide's mathematical model (BRI) and our changes", 1)
para("The guide's document proposed the Board Recoverability Index (BRI): it turns the detected defects into a damage "
     "score, a recoverable area and a grade A / B / C / Reject. It does not change detection accuracy; it is the "
     "quality-assessment part of the project. We implemented all six steps exactly as given, and our code reproduces "
     "the worked example of the document (damage score 0.020860, grade B).")
formula("s_i = w_c × conf_i × p_i,   p_i = 1 − exp(−d_i / λ),   DIS = Σ max_i s_i / A_board,   A_rec = A_board (1 − DIS)")
figure(FIG / "bri_geometry.png", 14, "Fig. Our images are sections of a board, with black conveyor on both sides; the trimmable edges are the board's sides.")
table([["Part of BRI", "Guide's version", "Our change and why"],
       ["Distance d_i", "to the nearest image border", "to the nearest side edge of the board (board mask): image borders include ~160 px of conveyor, and the top and bottom are cuts between sections, not edges a sawmill can trim"],
       ["Board area", "image width x height", "wood pixels only: the conveyor is ~12 % of each image"],
       ["Grade thresholds", "0.01 / 0.03 / 0.06 (illustrative)", "fitted on validation boards: these defaults grade 72 % of our boards A; the detector gets its own thresholds (guide's calibrate_thresholds) because model confidences are below 1"],
       ["Severity weights", "8-class placeholder table", "same values for the 4 classes kept: Crack 1.00, Dead_Knot 0.60, Live_Knot 0.40, Resin 0.30 (await the guide's ranking)"]],
      [3, 3.8, 9.7], size=9)
para("Validation (no expert grades yet): the reference grade is BRI computed on the annotators' boxes; the model grade "
     "is BRI on the detector's boxes; all thresholds are fitted on validation and agreement is measured on the 12 test boards.")
pi, pb = bri["test_per_image"], bri["test_per_board"]
ps = bri_single["test_per_image"]
table([["Agreement with the annotation-based grade", "Deployed detector, per image", "Deployed detector, per board", "Best single model, per image"],
       ["Exact grade", f"{pi['exact']*100:.1f} %", f"{pb['exact']*100:.1f} %", f"{ps['exact']*100:.1f} %"],
       ["Within one grade", f"{pi['within_one']*100:.1f} %", f"{pb['within_one']*100:.1f} %", f"{ps['within_one']*100:.1f} %"],
       ["Quadratic weighted kappa", f"{pi['qwk']:.2f}", f"{pb['qwk']:.2f}", f"{ps['qwk']:.2f}"],
       ["Mean abs. error of recoverable area", f"{pi['mae_recoverable_pct']:.2f} points", f"{pb['mae_recoverable_pct']:.2f} points", f"{ps['mae_recoverable_pct']:.2f} points"]],
      [5.6, 3.6, 3.6, 3.6], size=9)
para("The grading also feeds a value estimate in the application: value = board volume x recoverable share x grade-A "
     "price x grade factor (all inputs editable).")
para("", "Open points for the guide: ")
doc.paragraphs[-1].add_run("(1) approve the changes above; (2) rank the four defect classes by severity; (3) blind expert "
                           "grading of 30-50 board images (e.g. via the IWST contact) to calibrate the weights and thresholds "
                           "and report kappa against an expert, as in Section 5 of her document.")

# ======================================================================= 7 deliverables
doc.add_heading("7. Deliverables and suggested contributions for the paper", 1)
bullet("final presentation (33 slides) for 24 October; status update on BRI for the guide.", "Documents: ")
bullet("web application (default: the deployed fusion detector, with grading and value estimate).", "Application: ")
bullet("prepare_dataset, build_balanced, train_yolo11m, train_wiou, evaluate, eval_capture, math_model, bri, "
       "bri_validate, enhance_study; every decision and result logged in NOTES.md.", "Code: ")
para("Contributions the paper can claim:", bold_lead=None)
number("A leakage-free, board-level evaluation of YOLO11 on industrial line-scan timber images, showing that model size is not the bottleneck.")
number("A mathematical model (class-balancing rule + weighted box fusion) that raises mAP50 from 69.2 to 74.7 with a significant board-level confidence interval.")
number("An adapted Board Recoverability Index for line-scan geometry, validated with weighted kappa 0.83.")
number("Negative results reported with intervals: enhancement (CLAHE), Wise-IoU, per-class thresholds and higher resolution did not help.")

doc.save(str(OUT))
print("saved", OUT)
