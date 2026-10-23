"""
fill_report_pages.py - page numbers for the contents and the lists of figures / tables of the Phase-I report.
Word finds the page (as printed in the chapter section) of every heading / caption behind a {{PG:key}} token;
python-docx then replaces the tokens; Word exports the PDF.
"""
import re
import win32com.client
from docx import Document

PATH = r"C:\Users\ASUS\Desktop\Major_Project\docs\Major_Project_Phase-I_Report_Draft.docx"


def word_session(fn):
    w = win32com.client.DispatchEx("Word.Application"); w.Visible = False; w.DisplayAlerts = 0
    d = None
    try:
        d = w.Documents.Open(PATH, False, False, False)
        return fn(d)
    finally:
        if d is not None:
            d.Close(False)
        w.Quit()


def find_pages(d):
    d.Repaginate()
    start = d.Sections(d.Sections.Count).Range.Start
    keys = sorted(set(re.findall(r"\{\{PG:(.*?)\}\}", d.Content.Text)))
    pages = {}
    for k in keys:
        rng = d.Range(start, d.Content.End)
        f = rng.Find; f.ClearFormatting(); f.MatchCase = True
        pages[k] = str(rng.Information(1)) if f.Execute(k) else "?"
    return pages


pages = word_session(find_pages)
doc = Document(PATH)
paras = list(doc.paragraphs) + [p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
n = 0
for p in paras:
    for r in p.runs:
        if "{{PG:" in r.text:
            r.text = re.sub(r"\{\{PG:(.*?)\}\}", lambda m: pages.get(m.group(1), "?"), r.text); n += 1
doc.save(PATH)


def export(d):
    d.Repaginate(); d.ExportAsFixedFormat(PATH.replace(".docx", ".pdf"), 17)
    return d.ComputeStatistics(2)


print("tokens replaced", n, "missing", [k for k, v in pages.items() if v == "?"], "pages", word_session(export))
