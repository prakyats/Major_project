"""
deck_tools.py - small python-pptx helpers to grow the 50 % review deck into the final deck without losing its template:
clone a slide (shapes, pictures, background), move slides, rewrite bullet text keeping the template's paragraph and
run formatting, refill a table (rows and columns added by copying the template's own rows/cells), and replace a picture.
"""
import copy
from pptx.util import Emu, Pt
from pptx.oxml.ns import qn

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def clone_slide(prs, src):
    """Append a copy of slide `src` (same layout, shapes, pictures, background) and return it."""
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rId, rel in src.part.rels.items():
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        if rel.is_external:
            rid_map[rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rId] = new.part.relate_to(rel.target_part, rel.reltype)
    bg = src._element.cSld.find(qn("p:bg"))
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    for shp in src.shapes:
        el = copy.deepcopy(shp._element)
        for node in el.iter():
            for attr in list(node.attrib):
                if attr.startswith("{%s}" % R_NS) and node.attrib[attr] in rid_map:
                    node.attrib[attr] = rid_map[node.attrib[attr]]
        new.shapes._spTree.append(el)
    return new


def move_slide(prs, slide, index):
    lst = prs.slides._sldIdLst
    target = [s for s in list(lst) if prs.part.related_part(s.rId) is slide.part][0]
    lst.remove(target)
    lst.insert(index, target)


def delete_slide(prs, slide):
    lst = prs.slides._sldIdLst
    target = [s for s in list(lst) if prs.part.related_part(s.rId) is slide.part][0]
    prs.part.drop_rel(target.rId)
    lst.remove(target)


def shape_by_text(slide, needle):
    for sh in slide.shapes:
        if sh.has_text_frame and needle in sh.text_frame.text:
            return sh
    raise KeyError(needle)


def shape_by_id(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def _clear_runs(p):
    for r in list(p._p):
        if r.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p._p.remove(r)


def set_text(shape, text):
    """Replace a single-paragraph text (title, caption) keeping the first run's formatting."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    rpr = copy.deepcopy(p0.runs[0]._r.find(qn("a:rPr"))) if p0.runs else None
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    _clear_runs(p0)
    r = p0.add_run(); r.text = text
    if rpr is not None:
        r._r.insert(0, rpr)


def set_bullets(shape, items, size=None, template_index=None, keep_first=False):
    """items: list of str or (bold_lead, rest) or ("", rest, level). The first non-empty paragraph of the template
    gives pPr and run formatting; a bold lead-in run is made bold. keep_first keeps paragraph 0 (e.g. a heading)."""
    tf = shape.text_frame
    pars = [p for p in tf.paragraphs]
    start = 1 if keep_first else 0
    tmpl = None
    cands = pars[start:] if template_index is None else [pars[template_index]]
    for p in cands:
        if p.runs and p.text.strip():
            tmpl = p; break
    if tmpl is None:
        tmpl = pars[start]
    ppr = copy.deepcopy(tmpl._p.find(qn("a:pPr")))
    rpr_plain, rpr_bold = None, None
    for r in tmpl.runs:
        rp = r._r.find(qn("a:rPr"))
        if rp is not None and rp.get("b") == "1" and rpr_bold is None:
            rpr_bold = copy.deepcopy(rp)
        elif rpr_plain is None:
            rpr_plain = copy.deepcopy(rp) if rp is not None else None
    if rpr_plain is None and tmpl.runs:
        rp = tmpl.runs[-1]._r.find(qn("a:rPr")); rpr_plain = copy.deepcopy(rp) if rp is not None else None
    for p in pars[start:]:
        p._p.getparent().remove(p._p)
    txBody = tf._txBody
    for it in items:
        if isinstance(it, str):
            lead, rest, lvl = "", it, None
        elif len(it) == 2:
            lead, rest, lvl = it[0], it[1], None
        else:
            lead, rest, lvl = it
        p_el = txBody.add_p() if hasattr(txBody, "add_p") else None
        if p_el is None:
            from pptx.oxml import parse_xml
            p_el = parse_xml('<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>')
            txBody.append(p_el)
        if ppr is not None:
            pp = copy.deepcopy(ppr)
            if lvl is not None:
                pp.set("lvl", str(lvl))
            p_el.insert(0, pp)
        for txt, bold in ((lead, True), (rest, False)):
            if not txt:
                continue
            r = p_el.add_r(); r.text = txt
            base = rpr_bold if (bold and rpr_bold is not None) else rpr_plain
            if base is not None:
                rp = copy.deepcopy(base)
                if bold:
                    rp.set("b", "1")
                elif rp.get("b") == "1":
                    rp.set("b", "0")
                if size:
                    rp.set("sz", str(int(size * 100)))
                r.insert(0, rp)
            elif size or bold:
                rp = r.get_or_add_rPr()
                if size:
                    rp.set("sz", str(int(size * 100)))
                if bold:
                    rp.set("b", "1")


def _set_cell(cell, text, bold=None, size=None):
    tf = cell.text_frame
    p0 = tf.paragraphs[0]
    rpr = None
    for p in tf.paragraphs:
        if p.runs:
            rp = p.runs[0]._r.find(qn("a:rPr")); rpr = copy.deepcopy(rp) if rp is not None else None; break
    if rpr is None:
        epr = p0._p.find(qn("a:endParaRPr"))
        if epr is not None:
            rpr = copy.deepcopy(epr); rpr.tag = qn("a:rPr")
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    _clear_runs(p0)
    r = p0.add_run(); r.text = str(text)
    if rpr is not None:
        r._r.insert(0, rpr)
    rp = r._r.get_or_add_rPr()
    if bold is not None:
        rp.set("b", "1" if bold else "0")
    if size:
        rp.set("sz", str(int(size * 100)))


def fill_table(table_shape, rows, col_widths=None, size=None, bold_rows=(), bold_first_col=False):
    """rows[0] = header. Rows/columns are added by copying the template's last body row / last column."""
    tbl = table_shape.table
    tbl_el = tbl._tbl
    nrow, ncol = len(rows), max(len(r) for r in rows)
    grid = tbl_el.find(qn("a:tblGrid"))
    gcols = grid.findall(qn("a:gridCol"))
    while len(gcols) < ncol:
        grid.append(copy.deepcopy(gcols[-1]))
        for tr in tbl_el.findall(qn("a:tr")):
            tr.append(copy.deepcopy(tr.findall(qn("a:tc"))[-1]))
        gcols = grid.findall(qn("a:gridCol"))
    while len(gcols) > ncol:
        grid.remove(gcols[-1])
        for tr in tbl_el.findall(qn("a:tr")):
            tr.remove(tr.findall(qn("a:tc"))[-1])
        gcols = grid.findall(qn("a:gridCol"))
    trs = tbl_el.findall(qn("a:tr"))
    while len(trs) < nrow:
        src = trs[-1] if len(trs) > 1 else trs[0]
        tbl_el.append(copy.deepcopy(src)); trs = tbl_el.findall(qn("a:tr"))
    while len(trs) > nrow:
        tbl_el.remove(trs[-1]); trs = tbl_el.findall(qn("a:tr"))
    total = table_shape.width
    widths = col_widths or [1.0 / ncol] * ncol
    s = float(sum(widths))
    for gc, w in zip(gcols, widths):
        gc.set("w", str(int(total * w / s)))
    for i, row in enumerate(rows):
        for j in range(ncol):
            txt = row[j] if j < len(row) else ""
            b = True if (i in bold_rows or (bold_first_col and j == 0 and i > 0)) else None
            _set_cell(tbl.cell(i, j), txt, bold=b, size=size)
    return tbl


def set_row_heights(table_shape, h_emu):
    for tr in table_shape.table._tbl.findall(qn("a:tr")):
        tr.set("h", str(int(h_emu)))


def replace_picture(slide, pic_shape, image_path, fit=True, box=None):
    """Remove pic_shape and place image_path inside the same box (or `box` = (left, top, w, h)), aspect kept, centred."""
    from PIL import Image
    l, t, w, h = box if box else (pic_shape.left, pic_shape.top, pic_shape.width, pic_shape.height)
    if pic_shape is not None:
        pic_shape._element.getparent().remove(pic_shape._element)
    iw, ih = Image.open(image_path).size
    if fit:
        sc = min(w / iw, h / ih)
        nw, nh = int(iw * sc), int(ih * sc)
        l, t, w, h = l + (w - nw) // 2, t + (h - nh) // 2, nw, nh
    return slide.shapes.add_picture(str(image_path), Emu(l), Emu(t), Emu(w), Emu(h))


def add_picture(slide, image_path, box):
    return replace_picture(slide, None, image_path, True, box)


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def set_text_first(shape, text):
    """Replace only the first paragraph's text (e.g. a heading inside a body text box)."""
    p0 = shape.text_frame.paragraphs[0]
    runs = p0.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
